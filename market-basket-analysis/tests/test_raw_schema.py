"""Raw Schema Pre-Flight Tests.

Validates that data/raw/Online Retail.xlsx exists, loads sheet 'Online Retail',
contains 541,909 data rows, and includes all 8 required columns:
InvoiceNo, StockCode, Description, Quantity, InvoiceDate, UnitPrice, CustomerID, Country.
"""

from pathlib import Path
import openpyxl
import pytest

RAW_DATA_PATH = Path("data/raw/Online Retail.xlsx")
REQUIRED_COLUMNS = [
    "InvoiceNo",
    "StockCode",
    "Description",
    "Quantity",
    "InvoiceDate",
    "UnitPrice",
    "CustomerID",
    "Country",
]
EXPECTED_DATA_ROWS = 541909
EXPECTED_SHEET_NAME = "Online Retail"


def test_raw_file_exists():
    """Verify raw Excel file exists at expected path."""
    assert RAW_DATA_PATH.exists(), f"Raw file not found at {RAW_DATA_PATH}"
    assert RAW_DATA_PATH.is_file(), f"{RAW_DATA_PATH} is not a file"


def test_sheet_name_and_header_schema():
    """Verify 'Online Retail' sheet exists and contains all required columns."""
    wb = openpyxl.load_workbook(str(RAW_DATA_PATH), read_only=True)
    assert EXPECTED_SHEET_NAME in wb.sheetnames, (
        f"Sheet '{EXPECTED_SHEET_NAME}' not found in {wb.sheetnames}"
    )

    sheet = wb[EXPECTED_SHEET_NAME]
    first_row = next(sheet.iter_rows(values_only=True))
    wb.close()

    assert list(first_row) == REQUIRED_COLUMNS, (
        f"Columns mismatch: expected {REQUIRED_COLUMNS}, got {list(first_row)}"
    )


def test_raw_row_count():
    """Verify raw dataset contains exactly 541,909 data rows."""
    wb = openpyxl.load_workbook(str(RAW_DATA_PATH), read_only=True)
    sheet = wb[EXPECTED_SHEET_NAME]
    max_row = sheet.max_row
    wb.close()

    # max_row includes header row
    data_rows = max_row - 1
    assert data_rows == EXPECTED_DATA_ROWS, (
        f"Expected {EXPECTED_DATA_ROWS} data rows, found {data_rows}"
    )
